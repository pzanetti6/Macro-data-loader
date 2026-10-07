from __future__ import annotations

import io
import re
from datetime import date, datetime

import numpy as np
from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel

import pandas as pd
import streamlit as st

from macro_core import (CANONICAL_VARS, US_VARS, COUNTRY_NAMES, EA19, DataError,
                        fetch_country_panel, fetch_macro_dataset, fetch_us_dataset)
from excel_tools import create_workbook, update_workbook, workbook_sheet_names


SHOCK_SHEET_BY_FREQ = {"M": "monthly_freq", "Q": "quarterly_freq", "A": "annual_freq"}

def _shock_period_key(v, freq: str):
    """
    Normalize heterogeneous shock dates to the selected macro frequency.

    Important rule: slash-formatted dates such as 01/03/99 are interpreted
    day-first (DD/MM/YY), which is the convention used in the supplied European
    shock workbooks. Thus 01/03/99 -> 1999-03, not 1999-01.
    """
    if v is None or (isinstance(v, float) and np.isnan(v)):
        raise ValueError("blank date")

    def _from_timestamp(d):
        d = pd.Timestamp(d)
        if freq == "A":
            return f"{d.year:04d}"
        if freq == "Q":
            return f"{d.year:04d}q{((d.month - 1)//3)+1}"
        return f"{d.year:04d}-{d.month:02d}"

    # Excel/openpyxl dates arrive as datetime objects: never stringify/reparse them.
    if isinstance(v, (datetime, pd.Timestamp)):
        return _from_timestamp(v)

    # Excel serial date fallback.
    if isinstance(v, (int, float, np.integer, np.floating)) and not isinstance(v, bool):
        # Four-digit year cells are valid annual dates.
        if freq == "A" and 1900 <= int(v) <= 2100 and float(v).is_integer():
            return f"{int(v):04d}"
        d = pd.Timestamp("1899-12-30") + pd.to_timedelta(float(v), unit="D")
        return _from_timestamp(d)

    s = str(v).strip()
    if not s:
        raise ValueError("blank date")

    # Explicit annual representation.
    if freq == "A" and re.fullmatch(r"(?:19|20)\d{2}(?:\.0)?", s):
        return s[:4]

    # Explicit quarter representations: 2024Q1, 2024-Q1, Q1 2024, etc.
    if freq == "Q":
        compact = re.sub(r"[\s_\-/.]", "", s).upper()
        m = re.fullmatch(r"((?:19|20)\d{2})Q([1-4])", compact)
        if m:
            return f"{m.group(1)}q{m.group(2)}"
        m = re.fullmatch(r"Q([1-4])((?:19|20)\d{2})", compact)
        if m:
            return f"{m.group(2)}q{m.group(1)}"

    # Explicit year-month forms are unambiguous.
    ym = re.fullmatch(r"((?:19|20)\d{2})[-/._](0?[1-9]|1[0-2])", s)
    if ym:
        d = pd.Timestamp(year=int(ym.group(1)), month=int(ym.group(2)), day=1)
        return _from_timestamp(d)
    ym2 = re.fullmatch(r"((?:19|20)\d{2})(0[1-9]|1[0-2])", s)
    if ym2:
        d = pd.Timestamp(year=int(ym2.group(1)), month=int(ym2.group(2)), day=1)
        return _from_timestamp(d)

    # European slash/dash full dates: DD/MM/YYYY or DD/MM/YY.
    # This specifically fixes 01/mm/yy shock sheets.
    if re.fullmatch(r"\d{1,2}[/-]\d{1,2}[/-]\d{2,4}", s):
        try:
            return _from_timestamp(pd.to_datetime(s, dayfirst=True, errors="raise"))
        except Exception:
            pass

    # ISO-like full dates YYYY-MM-DD / YYYY/MM/DD.
    if re.fullmatch(r"(?:19|20)\d{2}[/-]\d{1,2}[/-]\d{1,2}", s):
        try:
            return _from_timestamp(pd.to_datetime(s, yearfirst=True, errors="raise"))
        except Exception:
            pass

    # Month-name and other textual dates. Prefer European day-first interpretation.
    try:
        return _from_timestamp(pd.to_datetime(s, dayfirst=True, errors="raise"))
    except Exception:
        raise ValueError(f"'{v}' cannot be interpreted as a { {'M':'monthly','Q':'quarterly','A':'annual'}[freq] } date")


def _macro_period_key(v, freq: str):
    """Return canonical output dates: YYYY-MM, YYYYqQ, or YYYY."""
    if isinstance(v, pd.Period):
        if freq == "A":
            return f"{v.year:04d}"
        if freq == "Q":
            return f"{v.year:04d}q{v.quarter}"
        return f"{v.year:04d}-{v.month:02d}"
    if isinstance(v, (datetime, pd.Timestamp)):
        if freq == "A":
            return f"{v.year:04d}"
        if freq == "Q":
            return f"{v.year:04d}q{((v.month-1)//3)+1}"
        return f"{v.year:04d}-{v.month:02d}"
    return _shock_period_key(v, freq)

def _canonicalize_dataset_dates(datasets, freq: str):
    """Force the final Excel date column into one Stata-friendly text convention."""
    out={}
    for geo, df in datasets.items():
        x=df.copy()
        date_col="date" if "date" in x.columns else ("period" if "period" in x.columns else x.columns[0])
        x[date_col]=x[date_col].map(lambda v:_macro_period_key(v,freq))
        out[geo]=x
    return out


def _read_shock_sheet(file_bytes: bytes, freq: str):
    """
    Read the requested frequency sheet. If it is absent but a meeting-level
    sheet exists, aggregate meeting shocks by SUM into the requested frequency.
    """
    wanted = SHOCK_SHEET_BY_FREQ[freq]
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)

    # Accept common sheet-name variants, case-insensitively.
    # Examples: monthly_freq, monthly, month; quarterly_freq, quarterly, quarter;
    # annual_freq, annual, yearly, year.
    aliases = {
        "M": ["monthly_freq", "monthly", "month", "monthly frequency", "monthly-frequency"],
        "Q": ["quarterly_freq", "quarterly", "quarter", "quarterly frequency", "quarterly-frequency"],
        "A": ["annual_freq", "annual", "yearly", "year", "annual frequency", "annual-frequency"],
    }
    def _norm_sheet_name(s):
        return re.sub(r"[^a-z0-9]+", "_", str(s).strip().lower()).strip("_")

    normalized_sheets = {_norm_sheet_name(s): s for s in wb.sheetnames}
    matched_sheet = None
    for alias in aliases[freq]:
        key = _norm_sheet_name(alias)
        if key in normalized_sheets:
            matched_sheet = normalized_sheets[key]
            break

    def _sheet_to_df(ws):
        rows = list(ws.iter_rows(values_only=True))
        if len(rows) < 2:
            raise DataError(f"Shock sheet '{ws.title}' is empty.")
        headers = [str(x).strip() if x is not None else "" for x in rows[0]]
        date_candidates = [
            i for i,h in enumerate(headers)
            if h.lower() in ("date","meetingdate","meeting_date","year","period","quarter","month")
        ]
        date_idx = date_candidates[0] if date_candidates else 0
        if not headers[date_idx]:
            headers[date_idx] = "date"
        data = pd.DataFrame(rows[1:], columns=headers).dropna(how="all")
        return data, headers[date_idx]

    if matched_sheet is not None:
        data, date_col = _sheet_to_df(wb[matched_sheet])
        return data, date_col, matched_sheet

    # Flexible fallback for raw meeting-level shock workbooks.
    meeting_aliases = {"meeting", "meetings", "meeting_freq", "meeting_level", "meeting_frequency"}
    meeting_names = [s for s in wb.sheetnames if _norm_sheet_name(s) in meeting_aliases]
    if not meeting_names:
        raise DataError(
            f"Shock workbook has no recognized { {'M':'monthly','Q':'quarterly','A':'annual'}[freq] } sheet "
            f"(accepted examples: {', '.join(aliases[freq])}) and no meeting-level sheet from which "
            "the requested frequency can be constructed."
        )

    meeting_name = meeting_names[0]
    raw, date_col = _sheet_to_df(wb[meeting_name])

    # Convert meeting dates flexibly. Excel/openpyxl normally gives datetime objects;
    # text dates are handled by pandas.
    parsed_dates = []
    bad_dates = []
    for v in raw[date_col].tolist():
        try:
            if isinstance(v, (datetime, pd.Timestamp)):
                d = pd.Timestamp(v)
            elif isinstance(v, (int, float, np.integer, np.floating)) and not isinstance(v, bool):
                # Excel serial date fallback.
                d = pd.Timestamp("1899-12-30") + pd.to_timedelta(float(v), unit="D")
            else:
                d = pd.to_datetime(str(v).strip(), errors="raise")
            parsed_dates.append(d)
        except Exception:
            parsed_dates.append(pd.NaT)
            bad_dates.append(str(v))
    if bad_dates:
        raise DataError(
            f"Could not interpret meeting dates in '{meeting_name}'. Examples: "
            + ", ".join(bad_dates[:3])
        )

    # Identify usable value columns. Metadata such as Year/Month/source/rates are not
    # treated as shocks by default; shock/surprise-named columns are preferred.
    candidate_cols = [c for c in raw.columns if c != date_col and str(c).strip()]
    preferred = [
        c for c in candidate_cols
        if any(k in str(c).lower() for k in ("shock","surprise"))
    ]
    value_cols = preferred if preferred else candidate_cols

    converted = {}
    for c in value_cols:
        vals = raw[c].map(_shock_numeric)
        # Keep a column only if every nonblank observation is numeric and at least one exists.
        bad = raw[c].notna() & vals.isna()
        if not bad.any() and vals.notna().any():
            converted[c] = vals.astype(float)
    if not converted:
        raise DataError(
            f"No numeric shock/surprise columns could be identified in meeting-level sheet '{meeting_name}'."
        )

    temp = pd.DataFrame({"_meeting_date": parsed_dates, **converted})
    if freq == "M":
        temp["date"] = temp["_meeting_date"].map(lambda d: f"{d.year:04d}-{d.month:02d}")
    elif freq == "Q":
        temp["date"] = temp["_meeting_date"].map(lambda d: f"{d.year:04d}q{((d.month-1)//3)+1}")
    else:
        temp["date"] = temp["_meeting_date"].map(lambda d: f"{d.year:04d}")

    # Monetary-policy meeting shocks aggregate by summation within the lower frequency.
    agg = temp.groupby("date", as_index=False)[list(converted.keys())].sum(min_count=1)
    return agg, "date", f"{meeting_name} → derived {wanted}"


def _shock_numeric(v):
    """Recover numeric shock values even when Excel mistakenly formats them as dates/times."""
    if v is None:
        return np.nan
    if isinstance(v, (datetime, pd.Timestamp)):
        return float(to_excel(v))
    if isinstance(v, (int, float, np.integer, np.floating)) and not isinstance(v, bool):
        return float(v)
    try:
        return float(str(v).strip())
    except Exception:
        return np.nan

def _validate_and_prepare_shocks(file_bytes: bytes, freq: str, value_columns, macro_dates):
    raw, date_col, sheet = _read_shock_sheet(file_bytes, freq)
    if not value_columns:
        raise DataError("Select at least one shock series from the uploaded shock workbook.")

    keys=[]
    bad_dates=[]
    for v in raw[date_col].tolist():
        try:
            keys.append(_shock_period_key(v, freq))
        except Exception as exc:
            keys.append(None)
            bad_dates.append(str(exc))
    if bad_dates:
        sample="; ".join(bad_dates[:3])
        raise DataError(f"Shock dates do not match the selected {freq} frequency. {sample}")

    out=pd.DataFrame({"_period_key":keys})
    if out["_period_key"].duplicated().any():
        dups=out.loc[out["_period_key"].duplicated(False),"_period_key"].unique()[:5]
        raise DataError(
            "Shock file contains duplicate periods after date normalization: "
            + ", ".join(map(str, dups))
            + ". Check whether the source really has multiple observations in the same period."
        )

    for c in value_columns:
        if c == date_col or c not in raw.columns:
            continue
        vals=raw[c].map(_shock_numeric)
        bad=raw[c].notna() & vals.isna()
        if bad.any():
            examples=", ".join(map(str, raw.loc[bad,c].head(3).tolist()))
            raise DataError(f"Shock column '{c}' contains non-numeric values: {examples}")
        out[c]=vals.astype(float)

    macro_keys=[_macro_period_key(v,freq) for v in macro_dates]
    macro_set=set(macro_keys)
    shock_set=set(out["_period_key"].dropna())
    outside=sorted(shock_set-macro_set)
    if outside:
        sample=", ".join(outside[:6])
        raise DataError(
            f"Shock dates do not align with the selected macro sample ({macro_keys[0]} to {macro_keys[-1]}). "
            f"Outside dates include: {sample}"
        )

    # Partial coverage is legitimate (e.g. a shock series ending before the macro data).
    aligned=pd.DataFrame({"_period_key":macro_keys}).merge(out,on="_period_key",how="left")

    # GPT/ChatGPT meeting shocks are defined as zero in periods with no meeting.
    # Fill only BETWEEN the first and last observed period of that uploaded
    # series; do not invent zeros before the series begins or after it ends.
    shock_period_order = {k:i for i,k in enumerate(macro_keys)}
    observed_keys = [k for k in out["_period_key"].dropna() if k in shock_period_order]
    if observed_keys:
        lo = min(shock_period_order[k] for k in observed_keys)
        hi = max(shock_period_order[k] for k in observed_keys)
        inside = aligned.index.to_series().between(lo, hi)
        for c in value_columns:
            cname = str(c).lower().replace("_","").replace("-","")
            if c in aligned.columns and ("gpt" in cname or "chatgpt" in cname):
                aligned.loc[inside & aligned[c].isna(), c] = 0.0

    coverage={c:int(aligned[c].notna().sum()) for c in value_columns if c in aligned.columns}
    return aligned, coverage, sheet

def _attach_shocks(datasets, shock_aligned, shock_columns, freq):
    result={}
    for geo,df in datasets.items():
        x=df.copy()
        date_col="date" if "date" in x.columns else ("period" if "period" in x.columns else x.columns[0])
        x["_period_key"]=x[date_col].map(lambda v:_macro_period_key(v,freq))
        add=shock_aligned[["_period_key"]+[c for c in shock_columns if c in shock_aligned.columns]]
        x=x.merge(add,on="_period_key",how="left").drop(columns="_period_key")
        result[geo]=x
    return result


def _run_dataset_diagnostics(df: pd.DataFrame, freq: str):
    """Detect missing values, suspicious zeros, duplicates and robust outliers."""
    issues={"missing":[],"zeros":[],"outliers":[],"duplicates":[]}
    if df is None or df.empty:
        return issues
    date_col="date" if "date" in df.columns else ("period" if "period" in df.columns else df.columns[0])
    dates=df[date_col].astype(str)
    dup=dates[dates.duplicated(keep=False)].unique().tolist()
    if dup:
        issues["duplicates"]=dup[:12]

    numeric_cols=[]
    for c in df.columns:
        if c==date_col or str(c).endswith("_source"): continue
        s=pd.to_numeric(df[c],errors="coerce")
        if s.notna().any(): numeric_cols.append((c,s))

    for c,s in numeric_cols:
        nmiss=int(s.isna().sum())
        if nmiss:
            miss_dates=dates[s.isna()].tolist()
            issues["missing"].append((c,nmiss,miss_dates[:4]))
        # Flag zeros only when zeros are unusual for that series; shock columns are exempt
        # because non-meeting zeros are economically meaningful.
        cname=str(c).lower()
        nz=s.dropna()
        zmask=s.eq(0)
        if zmask.any() and not any(k in cname for k in ("shock","gpt","rr_","jk_")):
            zero_share=float(zmask.sum()/max(s.notna().sum(),1))
            if zero_share < .20 and (nz.abs().median() if not nz.empty else 0) > 1e-10:
                issues["zeros"].append((c,int(zmask.sum()),dates[zmask].tolist()[:4]))
        # Robust MAD outlier rule; report only very extreme observations.
        valid=s.dropna()
        if len(valid)>=12:
            med=float(valid.median())
            mad=float((valid-med).abs().median())
            if mad>0:
                rz=0.6745*(s-med)/mad
                omask=rz.abs()>6
                if omask.any():
                    vals=[(dates.iloc[i],float(s.iloc[i]),float(rz.iloc[i])) for i in range(len(s)) if bool(omask.iloc[i])]
                    issues["outliers"].append((c,vals[:5]))
    return issues

def _show_diagnostics(datasets, freq):
    st.subheader("Data diagnostics")
    total=sum(len(v) for v in datasets.values())
    st.caption(f"Automated checks across {len(datasets)} dataset(s) and {total:,} rows.")
    any_issue=False
    for geo,df in datasets.items():
        d=_run_dataset_diagnostics(df,freq)
        if d["duplicates"]:
            any_issue=True
            st.error(f"{geo}: duplicate dates detected — {', '.join(d['duplicates'])}.")
        if d["missing"]:
            any_issue=True
            examples="; ".join(f"{c}: {n} missing ({', '.join(ds)})" for c,n,ds in d["missing"][:6])
            st.warning(f"{geo}: missing observations detected. {examples}")
        if d["zeros"]:
            any_issue=True
            examples="; ".join(f"{c}: {n} unusual zero(s) ({', '.join(ds)})" for c,n,ds in d["zeros"][:6])
            st.warning(f"{geo}: suspicious isolated zeros detected. {examples}")
        if d["outliers"]:
            any_issue=True
            parts=[]
            for c,vals in d["outliers"][:5]:
                vv=", ".join(f"{dt}={val:.3g} (robust z {z:.1f})" for dt,val,z in vals[:3])
                parts.append(f"{c}: {vv}")
            st.info(f"{geo}: potential extreme outliers — " + "; ".join(parts))
    if not any_issue:
        st.success("Diagnostics passed: no missing values, suspicious isolated zeros, duplicate dates or extreme robust outliers detected.")

st.set_page_config(page_title="Macro Data Loader", page_icon="🌐", layout="wide")

st.markdown("""
<style>
.block-container {max-width: 1180px; padding-top: 1.35rem; padding-bottom: 4rem;}
[data-testid="stSidebar"] {display:none;}
h1,h2,h3,p,div,span,label,button,input,select {
    font-family: Helvetica, "Helvetica Neue", Arial, sans-serif;
}
h1 {font-weight:650; letter-spacing:-0.045em; font-size:2.75rem; line-height:1.02;}
h2,h3 {font-weight:620; letter-spacing:-0.025em;}
div[data-testid="stMetric"] {background:#f8fafc; border:1px solid #e5e7eb; padding:12px; border-radius:14px;}
div[data-testid="stFileUploader"] {border:1px solid #e5e7eb; border-radius:14px; padding:8px;}
.hero-shell {background:#111827; border-radius:22px; padding:3.4rem 3.6rem 3.1rem; margin:.35rem 0 2rem; box-shadow:0 12px 34px rgba(15,23,42,.10);}
.hero-copy {max-width:900px;}
.hero-kicker {font-size:.76rem; font-weight:500; letter-spacing:.14em; text-transform:uppercase; color:rgba(255,255,255,.62); margin-bottom:.7rem;}
.hero-title {font-family:Helvetica, "Helvetica Neue", Arial, sans-serif; font-size:3.35rem; line-height:1; font-weight:500; letter-spacing:-.05em; color:#fff; margin:0 0 .9rem 0;}
.hero-sub {font-family:Helvetica, "Helvetica Neue", Arial, sans-serif; font-size:1.08rem; line-height:1.55; font-weight:300; color:rgba(255,255,255,.80); max-width:760px;}
.hero-pills {display:flex; gap:.5rem; flex-wrap:wrap; margin-top:1.35rem;}
.hero-pills span {font-size:.76rem; color:rgba(255,255,255,.78); background:rgba(255,255,255,.06); border:1px solid rgba(255,255,255,.16); border-radius:999px; padding:.32rem .68rem;}
</style>
""", unsafe_allow_html=True)

st.markdown("""
<div class="hero-shell"><div class="hero-copy">
  <div class="hero-kicker">Macroeconomic research utility</div>
  <div class="hero-title">Macro Data Loader</div>
  <div class="hero-sub">
    Download, harmonise and export research-ready macro data for the euro area, European countries and the United States.
  </div>
  <div class="hero-pills">
    <span>ECB</span><span>Eurostat</span><span>FRED</span><span>Excel output</span>
  </div>
</div></div>
""", unsafe_allow_html=True)

st.subheader("1 · Output")
c1,c2,c3=st.columns([1.15,1.15,1])
with c1:
    action=st.selectbox("Task",["Create a new workbook","Update an existing workbook"])
with c2:
    geography=st.selectbox("Dataset",["Euro area","Euro-area countries","United States"])
with c3:
    freq_label=st.selectbox("Frequency",["Monthly","Quarterly","Annual"])
freq={"Monthly":"M","Quarterly":"Q","Annual":"A"}[freq_label]

# Calendar controls. Defaults target the earliest broadly available period for
# the core variables rather than an arbitrary 1999 cutoff.
_default_start = date(1996, 1, 1) if geography != "United States" else date(1948, 1, 1)
_default_end = date.today()
d1,d2=st.columns(2)
with d1:
    start_date=st.date_input(
        "From", value=_default_start,
        min_value=date(1913,1,1), max_value=date.today(),
        format="DD/MM/YYYY",
    )
with d2:
    end_date=st.date_input(
        "To", value=_default_end,
        min_value=date(1913,1,1), max_value=date.today(),
        format="DD/MM/YYYY",
    )
start_year=start_date.year
end_year=end_date.year

uploaded=None
sheet_name=None
if action=="Update an existing workbook":
    uploaded=st.file_uploader("Workbook to update",type=["xlsx"],key="update_workbook")
    if uploaded is not None:
        try:
            sheets=workbook_sheet_names(uploaded.getvalue())
            sheet_name=st.selectbox("Sheet to update",sheets)
        except Exception as exc:
            st.error(f"Could not read workbook: {exc}")

st.subheader("2 · Variables")
eu_labels={
    "real_gdp":"Real GDP","real_gdp_growth":"Real GDP growth","log_real_gdp":"Log real GDP",
    "hicp_level":"HICP level","hicp_rate":"HICP inflation","unemployment_rate":"Unemployment rate",
    "ip":"Industrial production","log_ip":"Log industrial production","energy_index":"Energy HICP index",
    "estr":"€STR","dfr":"Deposit facility rate (DFR)","mro":"Main refinancing operations rate (MRO)",
    "aaa_1y":"EA AAA 1Y spot yield","aaa_10y":"EA AAA 10Y spot yield",
    "sovereign_1y":"Sovereign 1Y yield","sovereign_10y":"Sovereign 10Y yield",
}
us_labels={
    "real_gdp":"Real GDP","real_gdp_growth":"Real GDP growth","log_real_gdp":"Log real GDP",
    "cpi_level":"CPI level","inflation_rate":"CPI inflation","unemployment_rate":"Unemployment rate",
    "ip":"Industrial production","log_ip":"Log industrial production","energy_cpi":"Energy CPI",
    "fed_funds":"Effective federal funds rate","sofr":"SOFR",
    "treasury_1y":"US Treasury 1Y yield","treasury_10y":"US Treasury 10Y yield",
}
labels=us_labels if geography=="United States" else eu_labels
available_vars=US_VARS if geography=="United States" else CANONICAL_VARS
selected_vars=st.multiselect(
    "Series to include",
    available_vars,
    default=available_vars,
    format_func=lambda x: labels.get(x,x.replace("_"," ").title()),
)

monthly_gdp_method="Chow-Lin (industrial production)"
if freq=="M" and any(v in selected_vars for v in ("real_gdp","real_gdp_growth","log_real_gdp")):
    monthly_gdp_method=st.selectbox(
        "Monthly GDP treatment",
        ["Chow-Lin (industrial production)","Linear interpolation","Quarterly step (repeat within quarter)","Leave blank"],
        help="Quarterly GDP is the benchmark. Chow-Lin uses monthly industrial production as the indicator."
    )

countries=EA19
panel_format="One sheet per country"
if geography=="Euro-area countries":
    countries=st.multiselect(
        "Countries",EA19,default=EA19,
        format_func=lambda x:f"{x} — {COUNTRY_NAMES.get(x,x)}"
    )
    panel_format=st.selectbox("Country workbook layout",["One sheet per country","Long panel"])

if geography=="United States":
    st.caption("US sources: BEA, BLS, Federal Reserve/FRBNY and U.S. Treasury series distributed through FRED. The workbook records the source and series ID beside each variable.")
else:
    st.caption("European sources: ECB Data Portal and Eurostat. The app keeps source provenance beside each series.")
    st.caption("EA AAA yields are common fitted AAA curves. Sovereign 10Y is the broad EA convergence yield for EA output and the selected country's own convergence yield in country output.")


# Public version: monetary-policy shock upload/integration is intentionally omitted.
shock_uploads = []
shock_file_previews = []

run_label = "Create workbook" if action == "Create a new workbook" else "Update workbook"
run = st.button(run_label, type="primary", use_container_width=True)

# Keep the latest successful download available for plotting even after
# Streamlit reruns caused by changing the plot selector.
if "last_datasets" not in st.session_state:
    st.session_state["last_datasets"] = None
if "last_frequency" not in st.session_state:
    st.session_state["last_frequency"] = None
if "last_output_bytes" not in st.session_state:
    st.session_state["last_output_bytes"] = None
if "last_output_filename" not in st.session_state:
    st.session_state["last_output_filename"] = None
if "last_meta" not in st.session_state:
    st.session_state["last_meta"] = None
if "last_summary" not in st.session_state:
    st.session_state["last_summary"] = None
if "last_action" not in st.session_state:
    st.session_state["last_action"] = None

if run:
    if end_date < start_date:
        st.error("The 'To' date must be on or after the 'From' date.")
        st.stop()
    if not selected_vars:
        st.error("Select at least one macro variable.")
        st.stop()
    if geography == "Country-specific" and not countries:
        st.error("Select at least one country.")
        st.stop()
    if action == "Update an existing workbook" and uploaded is None:
        st.error("Upload a workbook first.")
        st.stop()

    prog = st.progress(0, text="Starting…")
    try:
        if geography == "Euro area":
            prog.progress(0.15, text="Comparing ECB and Eurostat series…")
            df, meta = fetch_macro_dataset(
                "EA_AUTO", freq, int(start_year), int(end_year), selected_vars, monthly_gdp_method
            )
            datasets = {"EA": df}
            prog.progress(0.80, text="Building output…")
        elif geography == "United States":
            prog.progress(0.15, text="Downloading official US series…")
            df, meta = fetch_us_dataset(
                freq, int(start_year), int(end_year), selected_vars, monthly_gdp_method
            )
            datasets = {"US": df}
            prog.progress(0.80, text="Building output…")
        else:
            datasets, meta = fetch_country_panel(
                countries, freq, int(start_year), int(end_year), selected_vars,
                monthly_gdp_method,
                progress=lambda p, txt: prog.progress(min(int(p * 80), 80), text=txt),
            )

        shock_coverage = {}
        shock_columns = []
        if shock_uploads:
            if len(shock_file_previews) != len(shock_uploads):
                raise DataError("At least one uploaded shock workbook could not be read. Fix the file error shown above.")
            _all_names = [str(c) for _,_,_,_,cols in shock_file_previews for c in cols]
            _dups = sorted({c for c in _all_names if _all_names.count(c) > 1})
            if _dups:
                raise DataError(
                    "Duplicate shock column names across uploaded workbooks: "
                    + ", ".join(_dups)
                    + ". Rename them so every imported shock series has a unique header."
                )

            _base_df = next(iter(datasets.values()))
            _base_date_col = "date" if "date" in _base_df.columns else ("period" if "period" in _base_df.columns else _base_df.columns[0])
            _macro_dates = _base_df[_base_date_col].tolist()

            # Build one combined shock table on the macro date grid. Every value
            # column from every uploaded workbook is imported automatically.
            _combined = pd.DataFrame({"_period_key": [_macro_period_key(v, freq) for v in _macro_dates]})
            for _sf, _preview, _date_col, _sheet, _value_cols in shock_file_previews:
                if not _value_cols:
                    continue
                _aligned, _coverage, _used_sheet = _validate_and_prepare_shocks(
                    _sf.getvalue(), freq, _value_cols, _macro_dates
                )
                _combined = _combined.merge(
                    _aligned[["_period_key"] + _value_cols],
                    on="_period_key",
                    how="left",
                    validate="one_to_one",
                )
                shock_columns.extend(_value_cols)
                for _c, _n in _coverage.items():
                    shock_coverage[f"{_sf.name} · {_c}"] = _n

            datasets = _attach_shocks(datasets, _combined, shock_columns, freq)

        # Standardize the final output date representation regardless of source/input format.
        datasets = _canonicalize_dataset_dates(datasets, freq)

        if action == "Create a new workbook":
            out = create_workbook(datasets, freq, panel_format=panel_format)
            _slug = "ea" if geography == "Euro area" else ("us" if geography == "United States" else "countries")
            filename = f"macro_{_slug}_{freq_label.lower()}.xlsx"
            summary = {"updated_cells": 0, "added_rows": sum(len(x) for x in datasets.values())}
        else:
            if geography == "Euro-area countries":
                st.error("Updating one existing sheet is designed for a single aggregate dataset. Create a new workbook for the multi-country panel.")
                st.stop()
            _update_key = "US" if geography == "United States" else "EA"
            out, summary = update_workbook(
                uploaded.getvalue(), sheet_name, datasets[_update_key], frequency=freq,
                append_new_periods=True,
                extra_columns=shock_columns if shock_uploads else None,
            )
            base = uploaded.name.rsplit(".", 1)[0]
            filename = f"{base}_updated.xlsx"

        prog.progress(100, text="Done")
        st.success("Finished successfully.")
        st.session_state["last_datasets"] = datasets
        st.session_state["last_frequency"] = freq
        st.session_state["last_output_bytes"] = out
        st.session_state["last_output_filename"] = filename
        st.session_state["last_meta"] = [m.__dict__ for m in meta]
        st.session_state["last_summary"] = summary
        st.session_state["last_action"] = action

        c1, c2, c3 = st.columns(3)
        latest = max((m.latest_period for m in meta if m.latest_period), default="—")
        c1.metric("Latest source observation", latest)
        c2.metric("Series downloaded", len(meta))
        c3.metric("Periods in output", sum(len(x) for x in datasets.values()))

        if geography == "Euro area":
            st.subheader("Preview")
            st.dataframe(datasets["EA"].tail(18), use_container_width=True, hide_index=True)

        _show_diagnostics(datasets, freq)

        st.subheader("Series information")
        meta_df = pd.DataFrame([m.__dict__ for m in meta])
        if not meta_df.empty:
            # Make construction/manipulation explicit, not only source provenance.
            _construction = {
                "real_gdp_growth": "Derived: log growth of real GDP; 12-month / 4-quarter / 1-year change according to output frequency.",
                "log_real_gdp": "Derived: natural logarithm of real GDP.",
                "hicp_rate": "Derived: 12-month / 4-quarter / 1-year log change of the all-items HICP index.",
                "inflation_rate": "Derived: 12-month / 4-quarter / 1-year log change of CPI.",
                "log_ip": "Derived: natural logarithm of the industrial-production index.",
                "real_gdp": (
                    "Direct quarterly/annual benchmark. For monthly output: "
                    + ("Chow-Lin temporal disaggregation using industrial production as indicator."
                       if "Chow-Lin" in monthly_gdp_method else
                       "Monthly temporal conversion: " + monthly_gdp_method + ".")
                    if freq == "M" else
                    "Direct benchmark series converted only to the selected output frequency when necessary."
                ),
                "dfr": "ECB policy-rate level. Last announced rate is carried forward between change dates, then averaged within the requested period.",
                "mro": "ECB policy-rate level. Last announced rate is carried forward between change dates, then averaged within the requested period.",
                "estr": "Direct €STR observations; higher-frequency observations are averaged within the requested period.",
                "aaa_1y": "Direct ECB fitted euro-area AAA government-bond 1Y spot yield; daily observations are period-averaged.",
                "aaa_10y": "Direct ECB fitted euro-area AAA government-bond 10Y spot yield; daily observations are period-averaged.",
                "sovereign_10y": "Direct ECB convergence-purpose government yield; EA aggregate for EA output, country-specific for country output.",
                "fed_funds": "Direct effective federal funds rate; monthly source or period average at lower frequency.",
                "sofr": "Direct SOFR; observations are averaged within the requested period.",
                "treasury_1y": "Direct US Treasury/Federal Reserve constant-maturity 1Y yield; period-averaged where needed.",
                "treasury_10y": "Direct US Treasury/Federal Reserve constant-maturity 10Y yield; period-averaged where needed.",
            }
            # Meta variable names are not perfectly uniform, so match normalized labels.
            def _how(row):
                v=str(row.get("variable","")).strip().lower().replace(" ","_")
                aliases={
                    "real_gdp":"real_gdp","real_gdp_growth":"real_gdp_growth",
                    "industrial_production":"ip","cpi":"cpi_level",
                    "unemployment":"unemployment_rate","hicp":"hicp_level",
                    "sovereign_10y_yield":"sovereign_10y",
                    "ea_sovereign/convergence_10y_yield":"sovereign_10y",
                    "country_sovereign_10y_yield":"sovereign_10y",
                }
                k=aliases.get(v,v)
                if k in _construction: return _construction[k]
                note=str(row.get("note","") or "").strip()
                return note if note else "Downloaded directly from the stated source; only frequency alignment/period averaging is applied when required."
            meta_df["construction / manipulation"] = meta_df.apply(_how,axis=1)
            st.dataframe(meta_df, use_container_width=True, hide_index=True)
        if action == "Update an existing workbook":
            with st.expander("Update details"):
                st.json(summary)

    except DataError as exc:
        prog.empty()
        st.error(f"Data download problem: {exc}")
        st.caption("The app tries both ECB and Eurostat where an equivalent series is mapped, then reports the detailed failure.")
    except Exception as exc:
        prog.empty()
        st.exception(exc)


# The generated workbook lives in session state, so ordinary Streamlit reruns
# (changing selectors, opening expanders, plotting, etc.) never make the download vanish.
if st.session_state.get("last_output_bytes") is not None:
    st.download_button(
        "⬇️ Download Excel",
        data=st.session_state["last_output_bytes"],
        file_name=st.session_state.get("last_output_filename") or "macro_data.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        type="primary",
        use_container_width=True,
        key="persistent_excel_download",
    )
    st.caption("The latest completed workbook stays available here until you create or update another one.")

st.divider()
@st.fragment
def _plot_area():
    st.subheader("Plot a time series")
    _plot_sets = st.session_state.get("last_datasets")
    if _plot_sets:
        _geos = list(_plot_sets.keys())
        _pc1, _pc2 = st.columns([1, 2])
        with _pc1:
            _plot_geo = st.selectbox("Geography", _geos, key="plot_geography")
        _pdf = _plot_sets[_plot_geo].copy()

        # Only numeric data columns; hide source/provenance columns.
        _plot_vars = [
            c for c in _pdf.columns
            if c not in ("period", "date", "country", "geo")
            and not str(c).endswith("_source")
            and pd.api.types.is_numeric_dtype(_pdf[c])
        ]
        with _pc2:
            _plot_var = st.selectbox(
                "Variable",
                _plot_vars,
                key="plot_variable",
                format_func=lambda x: labels.get(x, x.replace("_", " ").title()),
            ) if _plot_vars else None

        if _plot_var:
            _xcol = "period" if "period" in _pdf.columns else ("date" if "date" in _pdf.columns else None)
            if _xcol:
                _chart = _pdf[[_xcol, _plot_var]].dropna().copy()
                _chart[_xcol] = _chart[_xcol].astype(str)
                if not _chart.empty:
                    st.line_chart(_chart.set_index(_xcol)[_plot_var], use_container_width=True)
                    _first = _chart.iloc[0]
                    _last = _chart.iloc[-1]
                    st.caption(
                        f"{labels.get(_plot_var, _plot_var)} · "
                        f"{_first[_xcol]} to {_last[_xcol]} · {_plot_geo}"
                    )
                else:
                    st.info("The selected variable has no observations in the downloaded range.")
            else:
                st.info("No period/date column was found for plotting.")
    else:
        st.caption("Create or update a workbook first. The downloaded series will then be available here for plotting.")

_plot_area()

with st.expander("Variable definitions, construction and sources"):
    st.markdown("""
**Dates and availability**
- The **From / To** controls use calendar dates. Output is still aligned to the selected monthly, quarterly or annual frequency.
- The default start is **1996 for European datasets**, which is close to the beginning of the core harmonised EA macro series, and **1948 for the US**, when the main GDP/labour-market core is broadly available.
- Individual variables can begin later. The app never invents pre-history: unavailable observations remain blank.

**European macro variables**
- **Real GDP** — official quarterly real-GDP benchmark. Annual output is frequency-converted from the official benchmark. Monthly GDP is not directly observed: the selected method is used, with **Chow–Lin + industrial production** as the preferred default.
- **Real GDP growth** — derived from real GDP using the appropriate year-on-year log change.
- **Log real GDP** — natural logarithm of the real-GDP level.
- **HICP level** — directly downloaded all-items harmonised consumer-price index.
- **HICP inflation** — derived from the HICP level as the year-on-year log change.
- **Unemployment** — directly downloaded harmonised unemployment rate.
- **Industrial production** — directly downloaded official industrial-production index; frequency aggregation is applied only when required.
- **Log industrial production** — natural logarithm of industrial production.
- **Energy index** — directly downloaded HICP energy/special aggregate.
- **DFR / MRO** — official ECB policy-rate levels. The prevailing rate is **carried forward between policy changes**, then averaged within month/quarter/year, so unchanged periods are not missing.
- **€STR** — official €STR observations, averaged to the requested frequency where needed.
- **EA AAA 1Y / 10Y** — ECB fitted AAA government-bond spot curves; daily observations are period-averaged. Their official history starts later than the core macro data.
- **Sovereign 10Y** — ECB convergence-purpose government yield: broad EA aggregate for EA output and the individual country's yield for country output. It is not treated as equivalent to the AAA curve.

**United States**
- **Real GDP** — BEA real GDP (`GDPC1`) distributed through FRED. Monthly values are temporally disaggregated from the quarterly benchmark using the selected method.
- **CPI** — BLS CPI (`CPIAUCSL`) distributed through FRED.
- **CPI inflation** — derived year-on-year log change of CPI.
- **Unemployment** — BLS unemployment rate (`UNRATE`).
- **Industrial production** — Federal Reserve industrial production (`INDPRO`).
- **Energy CPI** — BLS energy CPI.
- **Federal funds rate** — effective federal funds rate.
- **SOFR** — Federal Reserve Bank of New York SOFR.
- **Treasury 1Y / 10Y** — US Treasury/Federal Reserve constant-maturity yields (`GS1`, `GS10`).

""")
